#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export the full list of members of an ArcGIS Enterprise / Portal.

Reimplemented from the original arcgis-based snippet (cod.txt) using only the
Python standard library (+ openpyxl for the Excel output), talking directly to
the Portal REST API:

    1. POST /sharing/rest/generateToken            -> obtain an access token
    2. GET  /sharing/rest/portals/self/users       -> list all members (paged)
    3. GET  /sharing/rest/community/users/{user}    -> per-member groups (opt.)

Output (same directory for both):
    - portal_users.csv
    - portal_users.xlsx
Columns: Username, LicenseType, FullName, Email, Role, Level, Disabled,
         Provider, Created, LastLogin, Modified, Groups
"""

import csv
import os
import ssl
import sys
import json
import getpass
import datetime
import urllib.parse
import urllib.request

# ---------------------------------------------------------------------------
# Default connection parameters
# ---------------------------------------------------------------------------
# These are only defaults: the portal URL, the username and the output folder
# are all requested interactively at run time. Press Enter at a prompt to keep
# the value shown in [square brackets].
DEFAULT_PORTAL_URL = "https://gis.ancpi.ro/portal"
DEFAULT_USERNAME = "ginu.popescu"
DEFAULT_OUTPUT_PATH = r"C:\#_LUCRU\Aplicatii\ExportPortalUsers"

# The password is never stored in this file. By default it is requested
# interactively at run time (input is hidden and kept only in memory).
# As an optional convenience, if the PORTAL_PASSWORD environment variable is
# set it is used instead of prompting.
PASSWORD_ENV_VAR = "PORTAL_PASSWORD"

# Base file name (used for both .csv and .xlsx) when a directory is given.
DEFAULT_FILE_STEM = "portal_users"


def prompt_with_default(label, default):
    """Ask the user for a value; empty input keeps the default shown."""
    try:
        answer = input("%s [%s]: " % (label, default)).lstrip("﻿").strip()
    except (EOFError, KeyboardInterrupt):
        raise RuntimeError("Input cancelled.")
    # Allow pasting a value wrapped in quotes.
    if len(answer) >= 2 and answer[0] == answer[-1] and answer[0] in "\"'":
        answer = answer[1:-1].strip()
    return answer or default


def get_password(uname):
    """Return the password: env var if present, otherwise an interactive prompt."""
    env_pw = os.environ.get(PASSWORD_ENV_VAR)
    if env_pw:
        return env_pw
    try:
        pw = getpass.getpass("Parola pentru %s: " % uname)
    except (EOFError, KeyboardInterrupt):
        raise RuntimeError("Password entry cancelled.")
    if not pw:
        raise RuntimeError("No password entered.")
    return pw

# Set to True to skip TLS certificate verification (self-signed / internal CA).
INSECURE_SSL = False

# Set to False to export ALL members; True keeps the original snippet's filter
# (exclude esri_* built-in accounts and keep only Viewer license types).
FILTER_LIKE_ORIGINAL = False

# Fetch each member's group memberships (one extra REST call per member).
# Set to False to skip it for a much faster run.
INCLUDE_GROUPS = True

# Write an Excel workbook in addition to the CSV (requires openpyxl).
WRITE_XLSX = True

PAGE_SIZE = 100  # members fetched per REST call

COLUMNS = [
    "Username", "LicenseType", "FullName", "Email", "Role", "Level",
    "Disabled", "Provider", "Created", "LastLogin", "Modified", "Groups",
]


def _make_ssl_context():
    if INSECURE_SSL:
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
        return ctx
    return None


def _post(url, params, ctx):
    data = urllib.parse.urlencode(params).encode("utf-8")
    req = urllib.request.Request(url, data=data, method="POST")
    with urllib.request.urlopen(req, context=ctx) as resp:
        return json.loads(resp.read().decode("utf-8"))


def _get(url, params, ctx):
    full = url + "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(full, method="GET")
    with urllib.request.urlopen(req, context=ctx) as resp:
        return json.loads(resp.read().decode("utf-8"))


def generate_token(base, uname, pw, ctx):
    """Request a short-lived REST token from the portal."""
    url = base + "/sharing/rest/generateToken"
    params = {
        "username": uname,
        "password": pw,
        "referer": base,
        "expiration": 60,  # minutes
        "f": "json",
    }
    result = _post(url, params, ctx)
    if "token" not in result:
        raise RuntimeError("Token request failed: %s" % json.dumps(result))
    return result["token"]


def fetch_all_users(base, token, ctx):
    """Page through /portals/self/users and return every member."""
    url = base + "/sharing/rest/portals/self/users"
    users = []
    start = 1
    while True:
        params = {
            "start": start,
            "num": PAGE_SIZE,
            "sortField": "username",
            "sortOrder": "asc",
            "token": token,
            "f": "json",
        }
        result = _get(url, params, ctx)
        if "error" in result:
            raise RuntimeError("User query failed: %s" % json.dumps(result))
        users.extend(result.get("users", []))
        next_start = result.get("nextStart", -1)
        if next_start is None or next_start <= 0:
            break
        start = next_start
    return users


def fetch_user_groups(base, token, uname, ctx):
    """Return a semicolon-separated list of group titles for one member."""
    url = base + "/sharing/rest/community/users/" + urllib.parse.quote(uname)
    try:
        result = _get(url, {"token": token, "f": "json"}, ctx)
    except Exception:
        return ""
    groups = result.get("groups") or []
    titles = [g.get("title", "") for g in groups if g.get("title")]
    return "; ".join(titles)


def keep_user(user):
    if not FILTER_LIKE_ORIGINAL:
        return True
    # Original snippet: exclude esri_* built-ins, keep only viewer licenses.
    uname = user.get("username", "")
    lic = user.get("userLicenseTypeId", "")
    return "esri_" not in uname and lic == "viewerUT"


def fmt_epoch(value):
    """Convert portal epoch-milliseconds to a readable local datetime string."""
    if value in (None, "", -1):
        return "Never" if value == -1 else ""
    try:
        ts = int(value) / 1000.0
        return datetime.datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, OSError, OverflowError):
        return str(value)


def build_row(user, groups):
    return [
        user.get("username", ""),
        user.get("userLicenseTypeId", ""),
        user.get("fullName", ""),
        user.get("email", ""),
        user.get("role", ""),
        user.get("level", ""),
        "Yes" if user.get("disabled") else "No",
        user.get("provider", ""),
        fmt_epoch(user.get("created")),
        fmt_epoch(user.get("lastLogin")),
        fmt_epoch(user.get("modified")),
        groups,
    ]


def resolve_output_stem(path):
    """Return the full path stem (without extension) for the output files."""
    if os.path.isdir(path) or path.endswith(("\\", "/")) or os.path.splitext(path)[1] == "":
        return os.path.join(path, DEFAULT_FILE_STEM)
    return os.path.splitext(path)[0]


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(COLUMNS)
        writer.writerows(rows)


def write_xlsx(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Portal Users"

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(row)

    # Auto-fit-ish column widths based on longest value (capped).
    for idx, header in enumerate(COLUMNS, start=1):
        longest = len(str(header))
        for row in rows:
            longest = max(longest, len(str(row[idx - 1])))
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, 60)

    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNS)), len(rows) + 1)
    wb.save(path)


def main():
    print("=" * 60)
    print("  Export Portal Users - ANCPI")
    print("=" * 60)
    print("Apasa Enter pentru a pastra valoarea implicita din [paranteze].")
    print()

    # --- Interactive parameters -----------------------------------------
    portal_url = prompt_with_default("URL portal GIS", DEFAULT_PORTAL_URL)
    uname = prompt_with_default("Utilizator (username)", DEFAULT_USERNAME)
    pw = get_password(uname)
    out_path = prompt_with_default(
        "Calea de salvare a fisierelor CSV si XLSX", DEFAULT_OUTPUT_PATH)
    print()

    base = portal_url.rstrip("/")
    ctx = _make_ssl_context()

    print("Connecting to %s ..." % base)
    token = generate_token(base, uname, pw, ctx)
    print("Token obtained. Fetching members ...")

    users = fetch_all_users(base, token, ctx)
    print("Portal returned %d member(s)." % len(users))

    users = [u for u in users if keep_user(u)]

    rows = []
    for i, user in enumerate(users, start=1):
        groups = ""
        if INCLUDE_GROUPS:
            groups = fetch_user_groups(base, token, user.get("username", ""), ctx)
            if i % 25 == 0 or i == len(users):
                print("  groups: %d/%d" % (i, len(users)))
        rows.append(build_row(user, groups))

    stem = os.path.abspath(resolve_output_stem(out_path))
    os.makedirs(os.path.dirname(stem) or ".", exist_ok=True)

    written_files = []

    csv_path = stem + ".csv"
    write_csv(csv_path, rows)
    written_files.append(csv_path)

    if WRITE_XLSX:
        try:
            xlsx_path = stem + ".xlsx"
            write_xlsx(xlsx_path, rows)
            written_files.append(xlsx_path)
        except ImportError:
            print("openpyxl not installed - skipped Excel output. "
                  "Install with: pip install openpyxl", file=sys.stderr)

    print()
    print("=" * 60)
    print("Gata. %d membri salvati in:" % len(rows))
    print("  Folder: %s" % os.path.dirname(stem))
    for fp in written_files:
        print("  - %s" % fp)
    print("=" * 60)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # noqa: BLE001 - top-level friendly error
        print("ERROR: %s" % exc, file=sys.stderr)
        sys.exit(1)
